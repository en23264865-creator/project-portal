from flask import Blueprint, send_file
from models import User, Project, Evaluation
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

export_bp = Blueprint('export', __name__)

HEADER_FILL  = PatternFill('solid', fgColor='1a1a2e')
HEADER_FONT  = Font(bold=True, color='FFFFFF', size=11)
SUBHDR_FILL  = PatternFill('solid', fgColor='5c49e0')
SUBHDR_FONT  = Font(bold=True, color='FFFFFF', size=10)
CENTER       = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT         = Alignment(horizontal='left',   vertical='center', wrap_text=True)
THIN_BORDER  = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD'),
)


def _title_row(ws, title, ncols):
    ws.append([title] + [''] * (ncols - 1))
    ws.merge_cells(f'A1:{get_column_letter(ncols)}1')
    cell = ws['A1']
    cell.font      = Font(bold=True, color='FFFFFF', size=13)
    cell.fill      = HEADER_FILL
    cell.alignment = CENTER
    ws.row_dimensions[1].height = 28


def _header_row(ws, headers, row=2):
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.font      = SUBHDR_FONT
        cell.fill      = SUBHDR_FILL
        cell.alignment = CENTER
        cell.border    = THIN_BORDER
    ws.row_dimensions[row].height = 20


def _set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _data_row(ws, values):
    ws.append(values)
    r = ws.max_row
    for col_idx in range(1, len(values) + 1):
        cell = ws.cell(row=r, column=col_idx)
        cell.alignment = LEFT
        cell.border    = THIN_BORDER
        # Alternate row shading
        if r % 2 == 0:
            cell.fill = PatternFill('solid', fgColor='F0EFFE')


# ─────────────────────────────────────────────────────────────
# STUDENTS EXPORT
# ─────────────────────────────────────────────────────────────
@export_bp.route('/students')
def export_students():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Students'

    headers = ['#', 'Student Name', 'Email Address', 'Contact Number',
               'Role', 'Registered On']
    _title_row(ws, '📋 Student List — Project Portal', len(headers))
    _header_row(ws, headers)

    students = User.query.filter_by(role='Student').order_by(User.name).all()
    for idx, s in enumerate(students, 1):
        _data_row(ws, [
            idx,
            s.name,
            s.email,
            s.contact or '—',
            s.role,
            s.created_at.strftime('%d %b %Y') if s.created_at else '—',
        ])

    _set_col_widths(ws, [5, 22, 30, 16, 12, 14])
    ws.freeze_panes = 'A3'

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='students.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─────────────────────────────────────────────────────────────
# PROJECTS EXPORT
# ─────────────────────────────────────────────────────────────
@export_bp.route('/projects')
def export_projects():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Projects'

    headers = ['#', 'Project Title', 'Domain', 'Type', 'Sem', 'Year',
               'Group Members', 'Lead Student Email', 'Guide Name', 'Guide Email',
               'Status', 'Evaluated?', 'Marks (/50)',
               'GitHub', 'Demo Link', 'Report Link', 'Code Link']
    _title_row(ws, '📁 Project List — Project Portal', len(headers))
    _header_row(ws, headers)

    for idx, p in enumerate(Project.query.order_by(Project.created_at).all(), 1):
        student_email = p.student.email  if p.student else '—'
        guide_name    = p.guide.name     if p.guide   else '—'
        guide_email   = p.guide.email    if p.guide   else '—'
        _data_row(ws, [
            idx,
            p.title,
            p.domain or '—',
            p.project_type or '—',
            p.sem  or '—',
            p.year or '—',
            p.group_members or '—',
            student_email,
            guide_name,
            guide_email,
            p.status,
            'Yes' if p.is_evaluated else 'No',
            p.marks if p.marks is not None else '—',
            p.github_link  or '—',
            p.demo_link    or '—',
            p.report_link  or '—',
            p.code_link    or '—',
        ])

    _set_col_widths(ws, [4, 28, 18, 8, 5, 6, 35, 28, 20, 28, 12, 10, 10, 35, 35, 35, 35])
    ws.freeze_panes = 'A3'

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='projects.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─────────────────────────────────────────────────────────────
# GUIDES EXPORT
# ─────────────────────────────────────────────────────────────
@export_bp.route('/guides')
def export_guides():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Guides'

    headers = ['#', 'Guide Name', 'Email', 'Contact', 'Domain 1', 'Domain 2', 'Domain 3',
               'No. of Students', 'No. of Projects']
    _title_row(ws, '🧑‍🏫 Guide List — Project Portal', len(headers))
    _header_row(ws, headers)

    guides = User.query.filter_by(role='Guide').order_by(User.name).all()
    for idx, g in enumerate(guides, 1):
        from models import Project as P
        proj_count    = P.query.filter_by(guide_id=g.id).count()
        from models import db
        from sqlalchemy import distinct
        student_count = db.session.query(distinct(P.student_id)).filter_by(guide_id=g.id).count()
        _data_row(ws, [
            idx, g.name, g.email, g.contact or '—',
            g.domain1 or '—', g.domain2 or '—', g.domain3 or '—',
            student_count, proj_count,
        ])

    _set_col_widths(ws, [4, 22, 30, 14, 18, 18, 18, 14, 14])
    ws.freeze_panes = 'A3'

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='guides.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─────────────────────────────────────────────────────────────
# ALL USERS EXPORT
# ─────────────────────────────────────────────────────────────
@export_bp.route('/all-users')
def export_all_users():
    wb = Workbook()
    ws = wb.active
    ws.title = 'All Users'

    headers = ['#', 'Name', 'Email', 'Contact', 'Role',
               'Domains / Batch Info', 'Registered On']
    _title_row(ws, '👥 All Users — Project Portal', len(headers))
    _header_row(ws, headers)

    users = User.query.order_by(User.role, User.name).all()
    for idx, u in enumerate(users, 1):
        if u.role == 'Guide':
            extra = ', '.join(u.domains_list())
        elif u.role == 'External Examiner':
            extra = f'Sem {u.exam_sem} · {u.exam_year}'
        else:
            extra = '—'
        _data_row(ws, [
            idx, u.name, u.email, u.contact or '—', u.role,
            extra,
            u.created_at.strftime('%d %b %Y') if u.created_at else '—',
        ])

    _set_col_widths(ws, [4, 22, 30, 14, 18, 30, 14])
    ws.freeze_panes = 'A3'

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='all_users.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
